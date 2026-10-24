"""
Nexo Sísmico · Base de datos NSR-10
Une en un solo esquema:
  (1) los edificios diseñados por el Prof. T. F. Rojas Vallejo (OpenSees, DES, Bogotá lacustre 200),
      con las entradas que definió el equipo (config_plantas_profesor.py) y sus resultados (Resumen_resultados_profesor.xlsx);
  (2) los casos analizados con el módulo nexo_nsr10.py (p. ej. el edificio del curso, grupo 1).
Las rutas E.030 (Perú) NO se mezclan aquí.
Nota del profesor (2026-10-24): varios resultados son de carácter únicamente teórico; algunas configuraciones
en planta y luces no se resolverían con pórticos resistentes a momento en la práctica. Las plantas ausentes
(P9-4P y P9-10P) fallaron en el dimensionamiento del algoritmo de OpenSees.
"""
import json, csv, sys
from openpyxl import load_workbook
import config_plantas_profesor as cp
from nexo_nsr10 import analizar, registro_bd
from rutas import COD, SAL

COLS = ["id", "fuente", "norma", "programa", "estado_dato", "n_pisos", "hn_m", "Lx_m", "Ly_m",
        "luces_x_m", "luces_y_m", "n_col", "area_m2", "uso", "grupo_uso", "I", "R",
        "D_entrepiso_kNm2", "L_entrepiso_kNm2", "Aa", "Av", "Fa", "Fv", "W_total_kN",
        "Tx_s", "Ty_s", "Tz_s", "CuTa_s", "Sa_x_g", "Sa_y_g", "Vx_kN", "Vy_kN", "dx_pct", "dy_pct", "columnas"]


def filas_profesor(xlsx=COD / "Resumen_resultados_profesor.xlsx"):
    ws = load_workbook(xlsx, data_only=True).active
    res = {r[1]: r[2:] for r in ws.iter_rows(min_row=4, values_only=True) if r[1]}
    E = {e["id"]: e for e in cp.edificios()}
    filas = []
    for k, v in res.items():
        e = E[k]
        Tx, Ty, Tz, Sax, Say, Vx, Vy, dx, dy, CuTa = v[:10]
        filas.append({
            "id": k, "fuente": "Prof. T. F. Rojas Vallejo (UNAL) · algoritmo OpenSees", "norma": "NSR-10",
            "programa": "OpenSees", "estado_dato": "teórico (profesor): configuración no necesariamente viable en pórticos reales",
            "n_pisos": e["pisos"], "hn_m": e["pisos"] * cp.H_ENTREPISO,
            "Lx_m": sum(e["sx"]), "Ly_m": sum(e["sy"]), "luces_x_m": e["sx"], "luces_y_m": e["sy"],
            "n_col": (len(e["sx"]) + 1) * (len(e["sy"]) + 1), "area_m2": sum(e["sx"]) * sum(e["sy"]),
            "uso": e["uso"], "grupo_uso": e["grupo"], "I": e["I"], "R": 7.0,
            "D_entrepiso_kNm2": e["D_ent"], "L_entrepiso_kNm2": e["L"],
            "Aa": None, "Av": None, "Fa": None, "Fv": None, "W_total_kN": round(Vx / Sax, 1) if Sax else None,
            "Tx_s": Tx, "Ty_s": Ty, "Tz_s": Tz, "CuTa_s": CuTa, "Sa_x_g": Sax, "Sa_y_g": Say,
            "Vx_kN": Vx, "Vy_kN": Vy, "dx_pct": dx, "dy_pct": dy,
            "columnas": "pendiente: transcribir de las imágenes del profesor",
        })
    return filas


def fila_modulo(ruta_json):
    datos = json.load(open(ruta_json, encoding="utf-8"))
    r = analizar(datos)
    if r["estado"] != "ok":
        raise ValueError(f"{ruta_json}: {r['estado']} {r['faltan'] or r['errores']}")
    f = registro_bd(datos, r)
    cols = datos["geometria"].get("columnas") or []
    f.update({"n_col": sum(c["n"] for c in cols) or None, "area_m2": None, "uso": "Oficinas",
              "grupo_uso": datos["sitio"]["grupo_uso"], "Tz_s": None,
              "D_entrepiso_kNm2": datos.get("D_entrepiso_kNm2"), "L_entrepiso_kNm2": datos.get("L_entrepiso_kNm2")})
    return f


def construir(salida=SAL / "base_datos_NSR10.csv", casos=(COD / "casos/caso1_completo.json",)):
    filas = filas_profesor() + [fila_modulo(c) for c in casos]
    with open(salida, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=COLS, extrasaction="ignore")
        w.writeheader()
        for f in filas:
            w.writerow({k: (json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v) for k, v in f.items()})
    return filas


if __name__ == "__main__":
    filas = construir()
    print(f"{len(filas)} registros → {SAL / 'base_datos_NSR10.csv'}")
